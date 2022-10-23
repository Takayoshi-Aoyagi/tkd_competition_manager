import os

from docx import Document

from excel_io import ResultsExcelReader


def replace_text(p, old_text, new_text):
    """ Replace placeholder text keeping styles """
    inline = p.runs
    # Loop added to work with runs (strings with same style)
    for i in range(len(inline)):
        if old_text in inline[i].text:
            text = inline[i].text.replace(old_text, new_text)
            inline[i].text = text


def generate_certificate(result, outdir):
    """ Read template word file, then replace placeholders. """
    tmpl_doc = Document('templates/tmpl_certificate.docx')

    for p in tmpl_doc.paragraphs:
        text = p.text
        if text == '%%EVENT%%':
            replace_text(p, text, result.event)
        elif text == '%%RANK%%':
            replace_text(p, text, result.rank.split('.')[0])
        elif text == '%%CLASSIFICATION%%':
            replace_text(p, text, result.classification)
        elif text == '%%NAME%%':
            replace_text(p, text, result.name)
    fname = f'{result.event}_{result.classification}_{result.rank}.docx'
    tmpl_doc.save(os.path.join(outdir, fname))


def main(result_file='winners.xlsx', outdir='certificates'):
    results = ResultsExcelReader(fpath=result_file).execute()
    for result in results:
        generate_certificate(result, outdir)


if __name__ == '__main__':
    main()
